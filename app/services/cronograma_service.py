from datetime import datetime
from openpyxl import load_workbook

def _clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = " ".join(value.split())
        return value or None
    return value

def parse_cronograma(path: str) -> dict:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    delegacion = _clean(ws["B8"].value)
    tipo_eleccion = _clean(ws["B9"].value)
    periodo_texto = _clean(ws["B10"].value)

    current = {
        "dia": None,
        "fecha": None,
        "provincia": None,
        "canton": None,
    }

    recorridos = []

    for row in range(13, ws.max_row + 1):
        funcionario = _clean(ws.cell(row, 1).value)
        dia = _clean(ws.cell(row, 2).value)
        fecha = ws.cell(row, 3).value
        hora = _clean(ws.cell(row, 4).value)
        provincia = _clean(ws.cell(row, 5).value)
        canton = _clean(ws.cell(row, 6).value)
        parroquia = _clean(ws.cell(row, 7).value)
        sector = _clean(ws.cell(row, 8).value)

        # Ignorar filas completamente vacías o filas ajenas a recorridos.
        if not any([funcionario, dia, fecha, hora, provincia, canton, parroquia, sector]):
            continue

        # Heredar valores cuando Excel usa una fila de continuación.
        if dia:
            current["dia"] = dia
        if fecha:
            if isinstance(fecha, datetime):
                current["fecha"] = fecha.date()
            else:
                current["fecha"] = fecha
        if provincia:
            current["provincia"] = provincia
        if canton:
            current["canton"] = canton

        # Si no existe hora o parroquia, probablemente no es una fila de recorrido.
        if not hora or not parroquia:
            continue

        recorridos.append({
            "funcionario": funcionario,
            "dia": current["dia"],
            "fecha": current["fecha"],
            "hora_programada": str(hora),
            "provincia": current["provincia"],
            "canton": current["canton"],
            "parroquia": parroquia,
            "sector": sector,
        })

    return {
        "delegacion": delegacion,
        "tipo_eleccion": tipo_eleccion,
        "periodo_texto": periodo_texto,
        "recorridos": recorridos,
    }
