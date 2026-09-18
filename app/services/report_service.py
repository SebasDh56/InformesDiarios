from pathlib import Path
from tempfile import TemporaryDirectory
from collections import defaultdict

from docx import Document
from docx.shared import Pt, Cm, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docxcompose.composer import Composer
from docxtpl import DocxTemplate, InlineImage


BASE_DIR = Path(__file__).resolve().parent.parent.parent

TEMPLATE_PRINCIPAL = (
    BASE_DIR
    / "templates_word"
    / "PLANTILLA_INFORME_ES2027_PICHINCHA.docx"
)

TEMPLATE_ANEXO = (
    BASE_DIR
    / "templates_word"
    / "PLANTILLA_ANEXO_ES2027_PICHINCHA.docx"
)


MESES = {
    1: "enero",
    2: "febrero",
    3: "marzo",
    4: "abril",
    5: "mayo",
    6: "junio",
    7: "julio",
    8: "agosto",
    9: "septiembre",
    10: "octubre",
    11: "noviembre",
    12: "diciembre",
}


DIAS = {
    0: "Lunes",
    1: "Martes",
    2: "Miércoles",
    3: "Jueves",
    4: "Viernes",
    5: "Sábado",
    6: "Domingo",
}


# ---------------------------------------------------------
# UTILIDADES
# ---------------------------------------------------------

def fecha_texto(fecha):
    return (
        f"{DIAS[fecha.weekday()]} "
        f"{fecha.day:02d} de "
        f"{MESES[fecha.month]} de "
        f"{fecha.year}"
    )


def hora_word(hora):
    """
    Convierte:
    09:30 -> 09h30
    13:00 -> 13h00
    """
    if not hora:
        return ""

    hora = str(hora).strip()

    if ":" in hora:
        h, m = hora.split(":", 1)
        return f"{h}h{m}"

    return hora


def separar_horario(horario):
    """
    '09:30 - 12:30'
    ->
    ('09h30', '12h30')
    """

    if not horario:
        return "", ""

    horario = str(horario)

    if " - " not in horario:
        return hora_word(horario), ""

    inicio, fin = horario.split(" - ", 1)

    return hora_word(inicio), hora_word(fin)


def texto_funcionarios(cronograma):

    funcionarios = []

    for rec in cronograma.recorridos:

        if (
            rec.funcionario
            and rec.funcionario not in funcionarios
        ):
            funcionarios.append(rec.funcionario)

    if not funcionarios:
        return ""

    if len(funcionarios) == 1:
        return funcionarios[0]

    return (
        ", ".join(funcionarios[:-1])
        + " y "
        + funcionarios[-1]
    )


def obtener_periodo(cronograma):

    fechas = [
        rec.fecha
        for rec in cronograma.recorridos
        if rec.fecha
    ]

    if not fechas:
        return ""

    inicial = min(fechas)
    final = max(fechas)

    if inicial.month == final.month:

        return (
            f"el {inicial.day:02d} al "
            f"{final.day:02d} de "
            f"{MESES[final.month]} de "
            f"{final.year}"
        )

    return (
        f"el {inicial.day:02d} de "
        f"{MESES[inicial.month]} de "
        f"{inicial.year} al "
        f"{final.day:02d} de "
        f"{MESES[final.month]} de "
        f"{final.year}"
    )


def parece_barrios(texto):

    if not texto:
        return False

    # El cronograma normalmente agrupa varios barrios
    # separados por coma.
    return "," in texto


# ---------------------------------------------------------
# PÁRRAFOS DEL CUERPO
# ---------------------------------------------------------

def configurar_parrafo(parrafo):

    pf = parrafo.paragraph_format

    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    # Ajustar aquí si el modelo Word tiene un valor diferente.
    pf.line_spacing = 1

    parrafo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def configurar_run(run, bold=False):

    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.bold = bold


def agregar_fecha(subdoc, fecha):

    p = subdoc.add_paragraph()

    configurar_parrafo(p)

    run = p.add_run(
        fecha_texto(fecha)
    )

    configurar_run(run, bold=True)

    return p


def agregar_intro_dia(
    subdoc,
    recorridos,
):

    if not recorridos:
        return

    canton = recorridos[0].canton or "Quito"

    if canton.upper() == "QUITO":

        texto = (
            "Recorrido realizado en el "
            "Distrito Metropolitano de Quito, "
            "en las siguientes parroquias:"
        )

    else:

        texto = (
            f"Recorrido realizado en el Cantón "
            f"{canton.title()}, "
            "en las siguientes parroquias:"
        )

    p = subdoc.add_paragraph()

    configurar_parrafo(p)

    run = p.add_run(texto)

    configurar_run(run)


def agregar_recorrido(
    subdoc,
    numero,
    recorrido,
):

    p = subdoc.add_paragraph()

    configurar_parrafo(p)

    # Sangría como la estructura utilizada
    # en los recorridos del modelo.
    p.paragraph_format.left_indent = Cm(0.7)
    p.paragraph_format.first_line_indent = Cm(-0.7)

    # Número
    run = p.add_run(f"{numero}.    ")

    configurar_run(run, bold=True)

    # Parroquia
    run = p.add_run(
        f"Parroquia {recorrido.parroquia or ''}"
    )

    configurar_run(run)

    # Barrio / sector
    if recorrido.sector:

        if parece_barrios(recorrido.sector):
            run = p.add_run(
                f", barrios: {recorrido.sector}"
            )
        else:
            run = p.add_run(
                f", sector {recorrido.sector}"
            )

        configurar_run(run)

    inicio, fin = separar_horario(
        recorrido.hora_programada
    )

    # Inicio
    if inicio:

        run = p.add_run(
            "; iniciando el recorrido a las "
        )

        configurar_run(run)

        run = p.add_run(inicio)

        configurar_run(run, bold=True)

        if recorrido.calle_inicial:

            run = p.add_run(
                f" en la calle "
                f"{recorrido.calle_inicial}"
            )

            configurar_run(run)

    # Final
    if fin:

        run = p.add_run(
            "; y, finalizando a las "
        )

        configurar_run(run)

        run = p.add_run(fin)

        configurar_run(run, bold=True)

        if recorrido.calle_final:

            run = p.add_run(
                f" en la calle "
                f"{recorrido.calle_final}"
            )

            configurar_run(run)

    run = p.add_run(".")

    configurar_run(run)


def agregar_nota_sin_articulos(subdoc):

    p = subdoc.add_paragraph()

    configurar_parrafo(p)

    run = p.add_run("Nota: ")

    configurar_run(run, bold=True)

    run = p.add_run(
        "Del recorrido realizado in situ, "
        "no se evidenciaron artículos promocionales."
    )

    configurar_run(run)


# ---------------------------------------------------------
# GENERACIÓN DEL CUERPO PRINCIPAL
# ---------------------------------------------------------

def generar_subdocumento_recorridos(
    template,
    cronograma,
):

    subdoc = template.new_subdoc()

    por_fecha = defaultdict(list)

    for rec in cronograma.recorridos:

        if rec.fecha:
            por_fecha[rec.fecha].append(rec)

    numero = 1

    for fecha in sorted(por_fecha):

        recorridos_dia = por_fecha[fecha]

        agregar_fecha(
            subdoc,
            fecha,
        )

        agregar_intro_dia(
            subdoc,
            recorridos_dia,
        )

        for rec in recorridos_dia:

            agregar_recorrido(
                subdoc,
                numero,
                rec,
            )

            numero += 1

        if not any(
            rec.hay_articulos
            for rec in recorridos_dia
        ):
            agregar_nota_sin_articulos(
                subdoc
            )

    return subdoc


# ---------------------------------------------------------
# ANEXOS
# ---------------------------------------------------------

def obtener_anexos(cronograma):

    anexos = []

    for recorrido in cronograma.recorridos:

        if not recorrido.hay_articulos:
            continue

        for anexo in recorrido.anexos:

            anexos.append(
                (recorrido, anexo)
            )

    return anexos


def generar_anexo(
    recorrido,
    anexo,
    numero,
    archivo_salida,
):

    doc = DocxTemplate(
        TEMPLATE_ANEXO
    )

    foto = ""

    if (
        anexo.foto_path
        and Path(anexo.foto_path).exists()
    ):

        foto = InlineImage(
            doc,
            anexo.foto_path,

            # Aproximadamente el ancho visual
            # utilizado en los anexos modelo.
            width=Mm(145),
        )

    contexto = {

        "numero_anexo":
            f"{numero:03d}",

        "provincia":
            recorrido.provincia or "",

        "canton":
            recorrido.canton or "",

        "parroquia":
            recorrido.parroquia or "",

        "circunscripcion":
            getattr(
                anexo,
                "circunscripcion",
                ""
            ) or "",

        "ciudad":
            getattr(
                anexo,
                "ciudad",
                ""
            ) or "",

        "sector":
            recorrido.sector or "",

        "zona":
            getattr(
                anexo,
                "zona",
                ""
            ) or "",

        "direccion":
            getattr(
                anexo,
                "direccion",
                None
            )
            or recorrido.calle_inicial
            or "",

        "referencia":
            getattr(
                anexo,
                "referencia",
                ""
            ) or "",

        "organizacion_politica":
            anexo.organizacion_politica
            or "",

        "dignidad":
            anexo.dignidad
            or "",

        "candidato":
            anexo.candidato
            or "",

        "leyenda":
            anexo.leyenda
            or "",

        "articulo":
            anexo.articulo
            or "",

        "medidas":
            anexo.medidas
            or "",

        "placa":
            anexo.placa
            or "",

        "hora":
            anexo.hora
            or "",

        "registro":
            anexo.registro
            or "",

        "cantidad":
            anexo.cantidad
            or 1,

        "evidencia_foto":
            foto,
    }

    doc.render(contexto)

    doc.save(
        archivo_salida
    )


# ---------------------------------------------------------
# INFORME COMPLETO
# ---------------------------------------------------------

def generar_word(
    cronograma,
    output_path,
    numero_informe,
):

    if not TEMPLATE_PRINCIPAL.exists():

        raise FileNotFoundError(
            "No se encontró la plantilla principal:\n"
            f"{TEMPLATE_PRINCIPAL}"
        )

    if not TEMPLATE_ANEXO.exists():

        raise FileNotFoundError(
            "No se encontró la plantilla de anexos:\n"
            f"{TEMPLATE_ANEXO}"
        )

    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    # --------------------------------
    # PLANTILLA PRINCIPAL
    # --------------------------------

    doc = DocxTemplate(
        TEMPLATE_PRINCIPAL
    )

    cuerpo = generar_subdocumento_recorridos(
        doc,
        cronograma,
    )

    funcionarios = texto_funcionarios(
        cronograma
    )

    periodo = obtener_periodo(
        cronograma
    )

    parrafo_equipo = (
        "La Delegación Provincial de Pichincha, "
        "en base a la normativa antes señalada, "
        "conformó un equipo de trabajo integrado "
        f"por los señores {funcionarios}, "
        "servidores de la Unidad Técnica Provincial "
        "de Fiscalización y Control del Gasto Electoral "
        "de este organismo provincial, quienes "
        "procedieron a realizar el recorrido de "
        "monitoreo de vías, de conformidad con el "
        "cronograma remitido a la Dirección Nacional "
        "de Fiscalización y Control del Gasto Electoral, "
        "para lo cual, me permito informar las novedades "
        f"registradas en el período comprendido entre {periodo}."
    )

    contexto = {

        "numero_informe":
            numero_informe,

        "parrafo_equipo":
            parrafo_equipo,

        "cuerpo_recorridos":
            cuerpo,

        # Estos después pueden salir
        # desde configuración/base de datos.

        "elaborado_por":
            "Lcdo. Franklin Portilla Hernández",

        "cargo_elaborado_por":
            "ANALISTA PROVINCIAL DE "
            "PARTICIPACION POLITICA 2",

        "revisado_por":
            "Ing. William Armando "
            "Narváez Lucero",

        "cargo_revisado_por":
            "Director Técnico Provincial de "
            "Participación Política de Pichincha",

        "aprobado_por":
            "Ab. Edwin Fabián Haro Aspiazu",

        "cargo_aprobado_por":
            "Director Delegación Provincial "
            "Electoral de Pichincha, Encargado",
    }

    # --------------------------------
    # RENDER PRINCIPAL
    # --------------------------------

    with TemporaryDirectory() as tmp:

        tmp_path = Path(tmp)

        principal_path = (
            tmp_path
            / "principal.docx"
        )

        doc.render(contexto)

        doc.save(
            principal_path
        )

        # --------------------------------
        # COMPOSITOR FINAL
        # --------------------------------

        principal_doc = Document(
            principal_path
        )

        composer = Composer(
            principal_doc
        )

        # --------------------------------
        # ANEXOS
        # --------------------------------

        anexos = obtener_anexos(
            cronograma
        )

        for numero, (
            recorrido,
            anexo
        ) in enumerate(
            anexos,
            start=1,
        ):

            anexo_path = (
                tmp_path
                / f"anexo_{numero:03d}.docx"
            )

            generar_anexo(
                recorrido,
                anexo,
                numero,
                anexo_path,
            )

            # Cada anexo empieza
            # en una nueva página.
            composer.doc.add_page_break()

            doc_anexo = Document(
                anexo_path
            )

            composer.append(
                doc_anexo
            )

        # --------------------------------
        # GUARDAR FINAL
        # --------------------------------

        composer.save(
            output_path
        )

    return output_path